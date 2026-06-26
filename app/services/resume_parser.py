import re
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MIN_CHARS_PER_PAGE = 80
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp", ".webp"}
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", *IMAGE_SUFFIXES}
ALLOWED_EXTENSIONS_DISPLAY = "PDF, Word (.docx), TXT, JPG, PNG"


@dataclass
class ParseResult:
    text: str
    method: str
    is_scanned: bool
    structure_note: str


def validate_resume_file(filename: str) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix == ".doc":
        raise ValueError(
            "Old Word format (.doc) is not supported. Please save as .docx or export to PDF and upload again."
        )
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"Unsupported file type '{suffix}'. HR can upload: {ALLOWED_EXTENSIONS_DISPLAY}."
        )


def parse_resume(file_bytes: bytes, filename: str) -> ParseResult:
    validate_resume_file(filename)
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(file_bytes)
    if suffix in {".docx"}:
        return _parse_docx(file_bytes)
    if suffix == ".txt":
        text = file_bytes.decode("utf-8", errors="ignore").strip()
        return ParseResult(
            text=_normalize_text(text),
            method="plain_text",
            is_scanned=False,
            structure_note="Plain text file.",
        )
    if suffix in IMAGE_SUFFIXES:
        return _parse_image(file_bytes)
    raise ValueError(f"Unsupported file type: {suffix}. Use {ALLOWED_EXTENSIONS_DISPLAY}.")


def parse_resume_text(file_bytes: bytes, filename: str) -> str:
    """Backward-compatible text-only entry point."""
    return parse_resume(file_bytes, filename).text


def extract_contact_info(text: str) -> dict[str, str]:
    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone_match = re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)
    name = _guess_name(text)
    return {
        "name": name,
        "email": email_match.group(0) if email_match else "",
        "phone": phone_match.group(0).strip() if phone_match else "",
    }


def normalize_unstructured_resume(
    text: str,
    parse_result: ParseResult,
    cancel_check: Callable[[], None] | None = None,
) -> str:
    """Use LLM to reorganize messy or OCR'd resumes into readable sections."""
    # Word/TXT already have clean text — never block on LLM
    if parse_result.method in ("docx", "plain_text"):
        return text
    if not parse_result.is_scanned and len(text) > 80:
        return text
    if not parse_result.is_scanned and len(text) > 200 and _has_clear_sections(text):
        return text

    from app.services.llm_client import llm_client

    if cancel_check:
        cancel_check()

    system = """You extract and reorganize resume content from messy, scanned, or unstructured text.
Preserve all factual details (names, dates, companies, skills, projects). Do not invent information.
Output clean plain text with sections: CONTACT, SUMMARY, EXPERIENCE, EDUCATION, PROJECTS, SKILLS.
If a section is missing, omit it."""
    user = f"""Source: {parse_result.method}. {parse_result.structure_note}

Raw resume text:
{text[:10000]}

Reorganize into structured plain text. No JSON."""
    try:
        content = llm_client.chat_text(system, user, fast=True)
        if len(content) > 100:
            return _normalize_text(content)
    except Exception:
        pass
    return text


def _parse_pdf(file_bytes: bytes) -> ParseResult:
    text = ""
    method = "pdf_text"

    try:
        import fitz

        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages = [page.get_text() for page in doc]
        doc.close()
        text = "\n".join(pages).strip()
        method = "pdf_pymupdf"
    except ImportError:
        pass

    if len(text) < MIN_CHARS_PER_PAGE:
        reader = PdfReader(BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        pypdf_text = "\n".join(pages).strip()
        if len(pypdf_text) > len(text):
            text = pypdf_text
            method = "pdf_pypdf"

    page_count = max(1, _pdf_page_count(file_bytes))
    chars_per_page = len(text) / page_count
    is_scanned = len(text) < MIN_CHARS_PER_PAGE * page_count or chars_per_page < MIN_CHARS_PER_PAGE

    if is_scanned or len(text) < MIN_CHARS_PER_PAGE:
        ocr_text = _ocr_pdf(file_bytes)
        if len(ocr_text) > len(text):
            text = ocr_text
            method = "pdf_ocr"
            is_scanned = True

    if not text:
        raise ValueError(
            "Could not extract text from PDF. If this is a scanned document, "
            "try uploading a clearer scan or a JPG/PNG image."
        )

    note = "Scanned or image-based PDF — OCR applied." if is_scanned else "Digital PDF text extraction."
    return ParseResult(
        text=_normalize_text(text),
        method=method,
        is_scanned=is_scanned,
        structure_note=note,
    )


def _parse_docx(file_bytes: bytes) -> ParseResult:
    doc = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                table_text.append(" | ".join(cells))
    parts = paragraphs + table_text
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("Could not extract text from DOCX")
    return ParseResult(
        text=_normalize_text(text),
        method="docx",
        is_scanned=False,
        structure_note="Word document — layout may vary.",
    )


def _parse_image(file_bytes: bytes) -> ParseResult:
    text = _ocr_image(file_bytes)
    if not text:
        raise ValueError("Could not read text from image. Ensure the scan is clear and upright.")
    return ParseResult(
        text=_normalize_text(text),
        method="image_ocr",
        is_scanned=True,
        structure_note="Scanned resume image — OCR applied.",
    )


def _ocr_pdf(file_bytes: bytes) -> str:
    try:
        from pdf2image import convert_from_bytes
        import pytesseract

        images = convert_from_bytes(file_bytes, dpi=200)
        parts = [pytesseract.image_to_string(img) for img in images]
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _ocr_image(file_bytes: bytes) -> str:
    try:
        from PIL import Image
        import pytesseract

        img = Image.open(BytesIO(file_bytes))
        return pytesseract.image_to_string(img).strip()
    except Exception:
        return ""


def _pdf_page_count(file_bytes: bytes) -> int:
    try:
        reader = PdfReader(BytesIO(file_bytes))
        return len(reader.pages)
    except Exception:
        return 1


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _has_clear_sections(text: str) -> bool:
    headers = re.findall(
        r"(?i)^(experience|education|skills|summary|work history|employment|projects)",
        text,
        re.MULTILINE,
    )
    return len(headers) >= 2


def _guess_name(text: str) -> str:
    for line in text.splitlines()[:8]:
        line = line.strip()
        if not line or "@" in line or re.search(r"\d{3}", line):
            continue
        if len(line) > 60:
            continue
        words = line.split()
        if 2 <= len(words) <= 5 and not any(w.lower() in {"resume", "curriculum", "cv"} for w in words):
            return line
    return "Unknown"
